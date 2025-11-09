import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
import os

def load_dataset_from_csv(file_path: str, train_ratio: float = None, test_ratio: float = None,
                          train_samples: int = None, test_samples: int = None):
    df = pd.read_csv(file_path)
    data = np.array(df)
    total_samples = len(data)
    i = 0
    while i < data.shape[1]:
        if np.isnan(data[:, i]).all():
            data = np.delete(data, i, axis=1)
        else:
            i += 1
    if train_ratio == 0.8 and test_ratio is None:
        return load_interspersed_data(data, total_samples)
    if train_samples or test_samples:
        if train_samples:
            test_samples = total_samples - train_samples
        else:
            train_samples = total_samples - test_samples
        train_data = data[:train_samples]
        test_data = data[train_samples:train_samples + test_samples]
    elif train_ratio or test_ratio:
        if train_ratio:
            test_ratio = 1 - train_ratio
        else:
            train_ratio = 1 - test_ratio
        train_count = int(train_ratio * total_samples)
        train_data = data[:train_count]
        test_data = data[train_count:]
    else:
        train_count = int(0.75 * total_samples)
        train_data = data[:train_count]
        test_data = data[train_count:]
    return train_data, test_data

def load_interspersed_data(data, total_samples, data_set_number=5):
    if data_set_number == 2:
        main_data = data[:600]
        extra_test_data = data[600:620]
        train_indices, test_indices = [], []
        for i in range(0, 600, 5):
            train_indices.extend(range(i, i + 4))
            test_indices.append(i + 4)
        train_data = main_data[train_indices]
        test_data_main = main_data[test_indices]
        test_data = np.vstack([test_data_main, extra_test_data])
    else:
        train_indices, test_indices = [], []
        for i in range(0, total_samples, 5):
            train_end = min(i + 4, total_samples)
            train_indices.extend(range(i, train_end))
            if i + 4 < total_samples:
                test_indices.append(i + 4)
        train_data = data[train_indices]
        test_data = data[test_indices]
    return train_data, test_data

def export_jc_iterations_to_excel(c_values, jc_values, iterations, filename="jc_iterations.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Jc_vs_Iterations"
    headers = ["Clusters (c)", "Objective Function (Jc)", "Iterations"]
    ws.append(headers)
    header_font = Font(bold=True)
    for col in range(1, len(headers) + 1):
        ws.cell(row=1, column=col).font = header_font
        ws.cell(row=1, column=col).alignment = Alignment(horizontal='center')
    for c, jc, iters in zip(c_values, jc_values, iterations):
        ws.append([c, jc, iters])
    wb.save(filename)
    print(f"[+] Exported Excel file → '{filename}' (no embedded charts)")

def export_cluster_file(data, cluster_ids, filename="C_output.csv"):
    x = data[:, 0]
    y = data[:, 1]
    rows = zip(x, y, cluster_ids)
    with open(filename, mode='w', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(["x", "y", "Cluster_ID"])
        writer.writerows(rows)
    print(f"[+] Exported cluster assignments to '{filename}' successfully!")

def plot_final_clusters(train_data, cluster_ids, num_clusters, centroids=None, title="Final Cluster Illustration"):
    import numpy as np
    import pandas as pd
    import matplotlib.pyplot as plt
    import os
    plt.figure(figsize=(8, 6))
    colors = plt.cm.tab10(np.linspace(0, 1, num_clusters))
    for k in range(num_clusters):
        cluster_points = train_data[cluster_ids == k]
        plt.scatter(cluster_points[:, 0], cluster_points[:, 1],
                    color=colors[k], label=f"Cluster {k+1}", s=25, alpha=0.8)
    try:
        if centroids is None:
            if os.path.exists("centroids_dev.csv"):
                df = pd.read_csv("centroids_dev.csv")
                centroids = df.select_dtypes(include=[np.number]).values
            else:
                raise FileNotFoundError("centroids_dev.csv not found.")
        if isinstance(centroids, (list, tuple)):
            centroids = np.array(centroids, dtype=float)
        elif isinstance(centroids, pd.DataFrame):
            centroids = centroids.values
        elif isinstance(centroids, np.ndarray):
            centroids = centroids.astype(float)
        if centroids.ndim == 1:
            centroids = centroids.reshape(-1, 2)
        plt.scatter(centroids[:, 0], centroids[:, 1],
                    color='black', marker='X', s=200,
                    edgecolors='white', linewidths=1.5, label='Centroids')
        for i, (x, y) in enumerate(centroids):
            plt.text(x + 0.1, y + 0.1, f"C{i+1}", fontsize=9, color='black', weight='bold')
    except Exception as e:
        print(f"[!] Could not plot centroids: {e}")
    plt.title(title)
    plt.xlabel("X")
    plt.ylabel("Y")
    plt.legend()
    plt.grid(True)
    os.makedirs("plots", exist_ok=True)
    safe_title = title.replace(" ", "_")
    output_path = os.path.join("plots", f"{safe_title}.png")
    plt.savefig(output_path, dpi=300, bbox_inches="tight")
    print(f"[+] Saved final cluster plot (with centroids) → '{output_path}'")
    plt.show()
